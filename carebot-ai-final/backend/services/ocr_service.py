"""Real local extraction; scanned PDF rendering uses PDFium (no Poppler setup)."""
from pathlib import Path
from io import BytesIO
import zipfile
import pytesseract
from PIL import Image, ImageOps, UnidentifiedImageError
from flask import current_app
from ..common import ApiError

Image.MAX_IMAGE_PIXELS=20_000_000

def image_ocr(im):
    if current_app.config['TESSERACT_CMD']: pytesseract.pytesseract.tesseract_cmd=current_app.config['TESSERACT_CMD']
    im=ImageOps.autocontrast(ImageOps.exif_transpose(im).convert('L'))
    im.thumbnail((3500,3500))
    try:
        words=pytesseract.image_to_data(im,config='--psm 3',output_type=pytesseract.Output.DICT,timeout=30)
    except pytesseract.TesseractNotFoundError: raise ApiError('Tesseract OCR is not installed or not on PATH. Set TESSERACT_CMD in the root .env.',503)
    except RuntimeError: raise ApiError('OCR timed out. Upload a smaller, clearer image.',422)
    lines={};conf=[];uncertain=[]
    for i,word in enumerate(words['text']):
        if not word.strip(): continue
        key=(words['block_num'][i],words['par_num'][i],words['line_num'][i])
        lines.setdefault(key,[]).append(word)
        c=float(words['conf'][i]);conf.append(c)
        if c<75: uncertain.append(word)
    return '\n'.join(' '.join(x) for x in lines.values()),conf,uncertain

def extract(path):
    ext=Path(path).suffix.lower();scores=[];uncertain=[];texts=[];method='text extraction'
    try:
        if ext in ('.png','.jpg','.jpeg','.webp'):
            with Image.open(path) as im: text,scores,uncertain=image_ocr(im)
            texts=[text];method='Tesseract OCR'
        elif ext=='.pdf':
            import pypdfium2 as pdfium
            with pdfium.PdfDocument(str(path)) as doc:
                if len(doc)>10: raise ApiError('Upload at most 10 PDF pages at a time.',422)
                for page in doc:
                    tp=page.get_textpage();text=tp.get_text_range();tp.close()
                    if len(text.strip())<20:
                        # Bound raster memory even for extremely large page dimensions.
                        w,h=page.get_size();scale=min(2.0,3500/max(w,h))
                        bitmap=page.render(scale=scale);im=bitmap.to_pil()
                        text,c,u=image_ocr(im);scores+=c;uncertain+=u;method='PDF text + Tesseract OCR'
                        im.close();bitmap.close()
                    texts.append(text);page.close()
        elif ext=='.docx':
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist())>25_000_000: raise ApiError('DOCX expands beyond the allowed size.',422)
            from docx import Document
            doc=Document(path);texts=[p.text for p in doc.paragraphs]
            texts += [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        elif ext in ('.txt','.csv'):
            texts=[Path(path).read_text(encoding='utf-8-sig')]
        else: raise ApiError('Unsupported report type.',415)
    except ApiError: raise
    except Exception: raise ApiError('The file is unreadable, corrupted, encrypted, or not the selected format.',422)
    text='\n'.join(texts).strip()
    if not text: raise ApiError('No readable text found. Upload a clear, upright image with readable text.',422)
    if len(text)>50000: raise ApiError('Document text is too long. Split the report into smaller files.',422)
    return {'text':text,'method':method,'mean_word_confidence':round(sum(scores)/len(scores),1) if scores else None,'uncertain_words':uncertain[:100],'review_required':True}
